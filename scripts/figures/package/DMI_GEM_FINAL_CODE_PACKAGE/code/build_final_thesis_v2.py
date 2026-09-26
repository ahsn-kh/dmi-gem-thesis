#!/usr/bin/env python3
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil

BASE_DOC = Path('/mnt/data/Bachelor_Thesis_26_COMPLETE_REVISED_FIGURES.docx')
OUT_DOC = Path('/mnt/data/Bachelor_Thesis_26_FINAL_FIGURES_EQUATIONS.docx')
FIG = Path('/mnt/data/final_thesis_build/final_v2_figures')
EQ = Path('/mnt/data/final_thesis_build/final_v2_equations')

shutil.copyfile(BASE_DOC, OUT_DOC)
doc = Document(str(OUT_DOC))
rels = doc.part.rels

def replace_blob(shape_index: int, image_path: Path):
    shape = doc.inline_shapes[shape_index]
    rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    image_part = rels[rid]._target
    image_part._blob = image_path.read_bytes()
    return shape

def set_width(shape, inches: float):
    shape.width = Inches(inches)

def set_natural_300dpi(shape, image_path: Path, max_width=5.72):
    im = Image.open(image_path)
    w = im.width / 300.0
    h = im.height / 300.0
    scale = min(1.0, max_width / w)
    shape.width = Inches(w * scale)
    shape.height = Inches(h * scale)

# User-requested figure replacements.
fig_replacements = {
    2: (FIG/'Figure_2_1_FINAL.png', 6.45),
    15: (FIG/'Figure_3_1_FINAL.png', 6.55),
    18: (FIG/'Figure_3_4_FINAL.png', 6.60),
    19: (FIG/'Figure_3_5_FINAL.png', 6.60),
    20: (FIG/'Figure_3_6_FINAL.png', 6.60),
}
for idx,(path,width) in fig_replacements.items():
    sh=replace_blob(idx,path); set_width(sh,width)

# Consistent equations: all rendered at the same 16-pt STIX math size and 300 dpi.
eq_replacements = {
    5: 'eq_2_1.png',
    6: 'eq_2_2.png',
    7: 'eq_2_3.png',
    8: 'eq_2_3b.png',
    10:'eq_2_4.png',
    11:'eq_2_5.png',
    12:'eq_2_6a.png',
    13:'eq_2_6b.png',
    14:'eq_2_7.png',
}
for idx,fn in eq_replacements.items():
    path=EQ/fn
    sh=replace_blob(idx,path)
    set_natural_300dpi(sh,path,max_width=5.72)

# Keep revised main figures centered. Paragraph indices are from the preserved base document.
for pidx in [72,149,167,173,180]:
    if pidx < len(doc.paragraphs):
        doc.paragraphs[pidx].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Save final document.
doc.save(str(OUT_DOC))
print(OUT_DOC)
